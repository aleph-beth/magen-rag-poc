from __future__ import annotations

import hashlib
import json
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pypdf import PdfReader

from magen_rag.models import Document

SUPPORTED = {".pdf", ".docx", ".html", ".htm", ".md", ".txt"}


def _hash_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if suffix == ".docx":
        doc = DocxDocument(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    raw = path.read_text(encoding="utf-8", errors="ignore")
    if suffix in {".html", ".htm"}:
        return BeautifulSoup(raw, "html.parser").get_text("\n")
    return raw


def ingest_folder(path: str) -> list[Document]:
    root = Path(path)
    docs: list[Document] = []
    for file in sorted(root.rglob("*")):
        if not file.is_file() or file.suffix.lower() not in SUPPORTED:
            continue
        text = _extract_text(file).strip()
        if not text:
            continue
        doc_id = _hash_text(str(file.relative_to(root)))[:12]
        docs.append(
            Document(
                doc_id=doc_id,
                source_path=str(file),
                doc_type=file.suffix.lower().lstrip("."),
                title=file.stem,
                text=text,
                hash_sha256=_hash_text(text),
            )
        )
    return docs


def write_manifest(documents: list[Document], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8") as fh:
        for doc in documents:
            payload = doc.model_dump(exclude={"text"}, mode="json")
            fh.write(json.dumps(payload, ensure_ascii=False) + "\n")
